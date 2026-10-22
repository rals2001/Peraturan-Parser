"""Streamlit app: Indonesian legal document (PDF/DOCX) -> structured review table (Word).

Single-file version. Deploy with:  app.py + requirements.txt
Sections: 1) patterns & hierarchy detector  2) PDF reader  3) DOCX reader
          4) table builder & Word export   5) Streamlit UI
"""
from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
from io import BytesIO

import pandas as pd
import pdfplumber
import regex as re
import streamlit as st
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


# ======================================================================
# 1. HIERARCHY DETECTOR
# ======================================================================
# ----------------------------------------------------------------- patterns
RE_BAB = re.compile(r"^BAB\s+([IVXLCDM]+)\s*[:.]?\s*(.*)$")
RE_BAGIAN = re.compile(r"^(Bagian|Paragraf)\s+(Ke\p{L}+|Pertama)\s*$")
RE_PASAL = re.compile(r"^Pasal\s+(\d+\s?[A-Z]?)\s*$")
RE_AYAT = re.compile(r"^\(\s*(\d{1,3}[a-z]?)\s*\)\s*(.*)$")
RE_HURUF = re.compile(r"^([a-z])[.)](?:\s+(.*))?$")
RE_ANGKA = re.compile(r"^(\d{1,3})\.(?:\s+(.*))?$")
RE_MENIMBANG = re.compile(r"^Menimbang\s*:?\s*(.*)$")
RE_MENGINGAT = re.compile(r"^Mengingat\s*:?\s*(.*)$")
RE_MENETAPKAN = re.compile(r"^Menetapkan\s*:?\s*(.*)$")
RE_MEMUTUSKAN = re.compile(r"^MEMUTUSKAN\s*:?\s*$")
RE_STOP = re.compile(r"^(PENJELASAN|LAMPIRAN)\b")
RE_PAGE_NUM = re.compile(
    r"^\s*(?:[-–—]\s*)?(?:hal(?:aman)?\.?\s*|page\s*)?\d{1,4}"
    r"\s*(?:[-–—]|/\s*\d+|of\s+\d+|dari\s+\d+)?\s*$", re.I)


def is_page_number(line: str) -> bool:
    """True for lines such as '3', '- 3 -', 'Halaman 3', '3/25', 'Page 3 of 25'."""
    return bool(RE_PAGE_NUM.match(line))


# --------------------------------------------------------------- data model
@dataclass
class Item:
    section: str                      # Menimbang | Mengingat | Menetapkan | Body
    bab: str = ""
    pasal: str = ""
    path: list = field(default_factory=list)   # [("ayat","(2)"),("huruf","a")]
    lines: list = field(default_factory=list)
    has_children: bool = False

    @property
    def text(self) -> str:
        return re.sub(r"\s+", " ", " ".join(self.lines)).strip()


@dataclass
class ParsedDocument:
    items: list = field(default_factory=list)
    bab_titles: dict = field(default_factory=dict)


# ------------------------------------------------------------------- parser
class HierarchyParser:
    def __init__(self):
        self.doc = ParsedDocument()
        self.mode = None          # None|menimbang|mengingat|menetapkan|body
        self.bab = self.pasal = ""
        self.stack: list = []
        self.cur: Item | None = None
        self.expect_title = None  # "bab" | "bagian" -> next line may be a title
        self.done = False

    # -- helpers
    def _new(self, section, path):
        it = Item(section, self.bab, self.pasal, list(path))
        self.doc.items.append(it)
        self.cur = it
        return it

    def _is_marker(self, ln):
        return any(r.match(ln) for r in (RE_AYAT, RE_HURUF, RE_ANGKA, RE_PASAL, RE_BAB))

    def _push(self, typ, val, text):
        if typ == "ayat":
            self.stack = [("ayat", f"({val})")]
        else:
            types = [t for t, _ in self.stack]
            if typ in types:                       # sibling -> replace, drop deeper levels
                self.stack = self.stack[: types.index(typ)] + [(typ, val)]
            else:                                  # first child of current level
                self.stack.append((typ, val))
        self._new("Body", self.stack)
        if text:
            self.cur.lines.append(text)

    # -- main entry per line
    def feed(self, ln: str):
        if self.done or not ln:
            return
        if self.mode in ("menetapkan", "body") and RE_STOP.match(ln):
            self.done = True
            return
        if RE_MEMUTUSKAN.match(ln):
            return
        # section keywords
        if self.mode is None and (m := RE_MENIMBANG.match(ln)):
            self.mode = "menimbang"; self.cur = None
            return self.feed(m.group(1).strip())
        if self.mode in (None, "menimbang") and (m := RE_MENGINGAT.match(ln)):
            self.mode = "mengingat"; self.cur = None
            return self.feed(m.group(1).strip())
        if self.mode in (None, "menimbang", "mengingat") and (m := RE_MENETAPKAN.match(ln)):
            self.mode = "menetapkan"; self._new("Menetapkan", [])
            return self.feed(m.group(1).strip())
        if self.mode is None:
            return                                  # preamble (title, "Dengan Rahmat...")

        # structural headings
        if self.mode in ("menetapkan", "body"):
            if m := RE_BAB.match(ln):
                self.mode = "body"; self.bab = f"BAB {m.group(1)}"
                self.pasal = ""; self.stack = []; self.cur = None
                self.doc.bab_titles[self.bab] = m.group(2).strip()
                self.expect_title = None if m.group(2) else "bab"
                return
            if RE_BAGIAN.match(ln):
                self.expect_title = "bagian"; self.cur = None
                return
            if m := RE_PASAL.match(ln):
                self.mode = "body"; self.pasal = f"Pasal {m.group(1).replace(' ', '')}"
                self.stack = []; self._new("Body", [])
                self.expect_title = None
                return
            if self.expect_title:                   # heading title line(s)
                if not self._is_marker(ln) and len(ln) < 120 and not ln.endswith((".", ":", ";")):
                    if self.expect_title == "bab":
                        t = self.doc.bab_titles.get(self.bab, "")
                        self.doc.bab_titles[self.bab] = (t + " " + ln).strip()
                    return
                self.expect_title = None

        # list markers
        if self.mode == "menimbang":
            if m := RE_HURUF.match(ln):
                self._new("Menimbang", [("huruf", m.group(1))]); self.cur.lines.append(m.group(2) or "")
                return
        elif self.mode == "mengingat":
            if m := RE_ANGKA.match(ln):
                self._new("Mengingat", [("angka", m.group(1))]); self.cur.lines.append(m.group(2) or "")
                return
        elif self.mode == "body":
            for rx, typ in ((RE_AYAT, "ayat"), (RE_HURUF, "huruf"), (RE_ANGKA, "angka")):
                if m := rx.match(ln):
                    return self._push(typ, m.group(1), (m.group(2) or "").strip())

        # continuation line -> merge into current provision
        if self.cur is None:
            self._new({"menimbang": "Menimbang", "mengingat": "Mengingat",
                       "menetapkan": "Menetapkan"}.get(self.mode, "Body"), [])
        self.cur.lines.append(ln)

    def finish(self) -> ParsedDocument:
        items = self.doc.items
        for i, it in enumerate(items[:-1]):         # mark parents (have deeper children)
            nxt = items[i + 1]
            n = len(it.path)
            if (nxt.section == it.section and nxt.pasal == it.pasal and nxt.bab == it.bab
                    and len(nxt.path) > n and nxt.path[:n] == it.path):
                it.has_children = True
        return self.doc


def detect_hierarchy(lines: list[str]) -> ParsedDocument:
    """Run the state machine over cleaned lines."""
    p = HierarchyParser()
    for ln in lines:
        ln = ln.strip()
        if is_page_number(ln):          # safety net: drop stray page numbers
            continue
        p.feed(ln)
    return p.finish()

# ======================================================================
# 2. PDF READER
# ======================================================================
def _norm(line: str) -> str:
    return re.sub(r"\d+", "#", line.lower()).strip()


def extract_pdf_lines(file) -> tuple[list[str], list[str]]:
    """Return (lines, warnings). `file` is a path or file-like object."""
    pages: list[list[str]] = []
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            txt = page.extract_text() or ""
            pages.append([l.strip() for l in txt.splitlines() if l.strip()])

    warnings = []
    if not any(pages):
        return [], ["PDF tidak memiliki teks (kemungkinan hasil scan). Gunakan OCR terlebih dahulu."]

    # Repeated lines at page edges (top/bottom 2 lines) = header/footer
    counter: Counter = Counter()
    for pl in pages:
        counter.update({_norm(l) for l in pl[:2] + pl[-2:]})
    threshold = max(3, int(len(pages) * 0.4))
    repeated = {k for k, v in counter.items() if v >= threshold} if len(pages) >= 3 else set()

    lines: list[str] = []
    for pl in pages:
        last = len(pl) - 1
        for i, l in enumerate(pl):
            at_edge = i < 2 or i > last - 2
            if at_edge and (_norm(l) in repeated or is_page_number(l)):
                continue
            lines.append(l)
    return lines, warnings

# ======================================================================
# 3. DOCX READER
# ======================================================================
def _iter_paragraphs(doc):
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, doc).rows:
                for cell in row.cells:
                    yield from cell.paragraphs


def extract_docx_lines(file) -> tuple[list[str], list[str]]:
    """Headers/footers live outside the body, so they are excluded automatically."""
    doc = Document(file)
    lines, auto_numbered = [], 0
    for p in _iter_paragraphs(doc):
        txt = p.text.strip()
        if not txt or is_page_number(txt):
            continue
        has_num = p._p.pPr is not None and p._p.pPr.numPr is not None
        if has_num and not any(r.match(txt) for r in (RE_AYAT, RE_HURUF, RE_ANGKA)):
            auto_numbered += 1
        lines.extend(l.strip() for l in txt.splitlines() if l.strip())
    warnings = []
    if auto_numbered:
        warnings.append(
            f"{auto_numbered} paragraf memakai penomoran otomatis Word (nomor tidak tersimpan sebagai teks). "
            "Ubah ke teks biasa (Select all > Home > Numbering off, atau simpan ke PDF) agar ayat/huruf/angka terdeteksi.")
    return lines, warnings

# ======================================================================
# 4. TABLE BUILDER & WORD
# ======================================================================
COLUMNS = ["No", "Bab", "Pasal", "Ayat", "Huruf", "Angka", "Bagian",
           "Draft RPM", "Masukan", "Justifikasi"]


def _label(item) -> str:
    if item.section == "Menimbang":
        return f"Menimbang Huruf {item.path[0][1]}" if item.path else "Menimbang"
    if item.section == "Mengingat":
        return f"Mengingat Poin {item.path[0][1]}" if item.path else "Mengingat"
    if item.section == "Menetapkan":
        return "Menetapkan"
    parts = [item.pasal] + [f"Ayat {v}" if t == "ayat" else f"{t} {v}" for t, v in item.path]
    return " ".join(p for p in parts if p)


def build_rows(doc: ParsedDocument) -> pd.DataFrame:
    """One row per leaf provision. Parents that only introduce sub-items
    (e.g. '(3) meliputi:') are skipped."""
    rows = []
    for it in doc.items:
        if not it.text or it.has_children:
            continue
        p = dict(it.path)
        rows.append({
            "Bab": it.bab if it.section == "Body" else "",
            "Pasal": it.pasal if it.section == "Body" else "",
            "Ayat": p.get("ayat", ""),
            "Huruf": p.get("huruf", ""),
            "Angka": p.get("angka", ""),
            "Bagian": _label(it),
            "Draft RPM": it.text,
            "Masukan": "", "Justifikasi": "",
        })
    df = pd.DataFrame(rows, columns=COLUMNS[1:])
    df.insert(0, "No", range(1, len(df) + 1))
    return df


def _shade(cell, color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def to_word_bytes(df: pd.DataFrame, title: str = "Tabel Review Draft RPM") -> bytes:
    """Landscape A4 Word document containing the review table."""
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.2)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(9)
    doc.add_heading(title, level=1)

    widths = [1.0, 1.8, 1.8, 1.3, 1.3, 1.3, 4.0, 8.2, 3.2, 3.2]   # cm, total ~27.1
    table = doc.add_table(rows=1, cols=len(COLUMNS))
    table.style = "Table Grid"
    table.autofit = False

    for i, name in enumerate(COLUMNS):                              # header row
        c = table.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(name)
        run.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
        _shade(c, "1F4E78")
    trPr = table.rows[0]._tr.get_or_add_trPr()                      # repeat header on each page
    hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true"); trPr.append(hdr)

    for rec in df.itertuples(index=False):
        cells = table.add_row().cells
        for i, v in enumerate(rec):
            cells[i].text = str(v)

    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ======================================================================
# 5. STREAMLIT UI  (input -> process -> download)
# ======================================================================
def main():
    st.set_page_config(page_title="Parser RPM", layout="centered")
    st.title("📑 Parser Draft Regulasi → Tabel Review (Word)")

    up = st.file_uploader("Upload dokumen (PDF atau DOCX)", type=["pdf", "docx"])
    if not up:
        st.info("Unggah file PDF atau DOCX untuk memulai.")
        return

    try:
        with st.spinner("Memproses dokumen..."):
            if up.name.lower().endswith(".pdf"):
                lines, warns = extract_pdf_lines(up)
            else:
                lines, warns = extract_docx_lines(up)
            df = build_rows(detect_hierarchy(lines))
            data = to_word_bytes(df) if not df.empty else None
    except Exception as e:  # noqa: BLE001
        st.error(f"Gagal memproses file: {e}")
        return

    for w in warns:
        st.warning(w)
    if data is None:
        st.error("Tidak ada struktur terdeteksi. Pastikan dokumen memuat Menimbang/Mengingat/Menetapkan/BAB/Pasal.")
        return

    st.success(f"Selesai: {len(df)} baris berhasil diekstrak.")
    st.download_button(
        "⬇️ Download Word", data,
        file_name=up.name.rsplit(".", 1)[0] + "_review.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


if __name__ == "__main__":
    main()
